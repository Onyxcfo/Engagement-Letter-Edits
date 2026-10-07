"""Build a client proposal deck from the reference proposal deck + a content JSON.

usage: python3 build_proposal.py content.json out.pptx --client "Client Name" [--template ref.pptx]
           [--logo logo.png] [--wordmark "LINE ONE|LINE TWO"] [--wordmark-plate "LINE|LINE|LINE"]
content.json = {"slots": [{"id": "...", "lines": ["..."]}, ...]}   (slot ids: see slots.md)
Without --logo, the client logo spots get a text wordmark built from --wordmark / --wordmark-plate.
"""
import sys, json, copy, re, os, argparse
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
NAVY, RED = RGBColor(0x1B, 0x2A, 0x49), RGBColor(0xAF, 0x1F, 0x29)
ap = argparse.ArgumentParser()
ap.add_argument('content'); ap.add_argument('out')
ap.add_argument('--client', required=True, help='client name for the footer and document title')
ap.add_argument('--template', default=os.path.join(HERE, '..', '..', '..', 'Onyx_AJAC_Proposal.pptx'))
ap.add_argument('--logo', default=None, help='client logo image; replaces the wordmark')
ap.add_argument('--wordmark', default=None, help='cover wordmark lines, separated by |')
ap.add_argument('--wordmark-plate', default=None, help='closing-slide wordmark lines, separated by |')
A = ap.parse_args()
GB, CLIENT, content, out = A.template, A.client, A.content, A.out
logo = A.logo if A.logo and os.path.exists(A.logo) else None
WM = (A.wordmark or CLIENT.upper()).split('|')
WMP = (A.wordmark_plate or A.wordmark or CLIENT.upper()).split('|')
C = {s['id']: s['lines'] for s in json.load(open(content))['slots']}
expected = set(json.load(open(os.path.join(HERE, 'slot_ids.json'))))
missing, unknown = expected - set(C), set(C) - expected
if unknown: print('WARNING unknown slot ids:', sorted(unknown))
if missing: print('WARNING missing slot ids (GB text kept):', sorted(missing))

p = Presentation(GB)
S = [list(s.shapes) for s in p.slides]
BOLD = re.compile(r'^\*\*(.+?)\*\*\s*(.*)$', re.S)

def fill_para(para, line):
    """Put `line` into paragraph `para`, keeping run formatting. **lead** = bold lead-in run."""
    runs = para.runs
    m = BOLD.match(line)
    if m:
        lead, rest = m.group(1) + ' ', m.group(2)
        if len(runs) >= 2:
            runs[0].text, runs[1].text = lead, rest
            for r in runs[2:]: r._r.getparent().remove(r._r)
        else:
            r0 = runs[0]; r1 = copy.deepcopy(r0._r); r0._r.addnext(r1)
            r0.text = lead; r0.font.bold = True
            para.runs[1].text = rest; para.runs[1].font.bold = None
    else:
        if not runs:
            return
        if len(runs) >= 2 and runs[0].font.bold and not runs[1].font.bold:
            # template had a bold lead-in but the new line has none: merge into the normal run
            runs[0]._r.getparent().remove(runs[0]._r); runs = para.runs
        runs[0].text = line
        for r in runs[1:]: r._r.getparent().remove(r._r)

def set_lines(tf, lines, spacer_idx=None, body_idx=None):
    """Rewrite a text frame to `lines`. spacer_idx/body_idx: GB paragraph indices used as
    templates when the frame interleaves spacer paragraphs (slide 6 company panel)."""
    paras = list(tf.paragraphs)
    templates = [copy.deepcopy(pp._p) for pp in paras]
    if spacer_idx is not None:
        spacer, body = templates[spacer_idx], templates[body_idx]
        seq = []
        for i, line in enumerate(lines):
            if i: seq.append(('spacer', ''))
            seq.append((0 if i == 0 else body_idx, line))
        new = []
        for t, line in seq:
            e = copy.deepcopy(spacer if t == 'spacer' else templates[t]); new.append((e, line))
    else:
        new = [(copy.deepcopy(templates[min(i, len(templates) - 1)]), line) for i, line in enumerate(lines)]
    txBody = paras[0]._p.getparent()
    for pp in paras: txBody.remove(pp._p)
    for e, line in new:
        txBody.append(e)
    for pp, (e, line) in zip(tf.paragraphs, new):
        if line != '':
            fill_para(pp, line)

def slot(si, idx, sid, **kw):
    if sid in C: set_lines(S[si - 1][idx].text_frame, C[sid], **kw)

def cell(si, idx, r, c, sid):
    if sid in C: set_lines(S[si - 1][idx].table.cell(r, c).text_frame, C[sid])

def replace_logo(si, idx, plate=None):
    """Swap the GB logo picture for the client logo file or a text wordmark."""
    pic = S[si - 1][idx]
    sl = p.slides[si - 1]
    left, top, w, h = pic.left, pic.top, pic.width, pic.height
    parent, pos = pic._element.getparent(), list(pic._element.getparent()).index(pic._element)
    parent.remove(pic._element)
    if logo:
        new = sl.shapes.add_picture(logo, left, top, height=h)
        if new.width > w:  # keep inside the old box
            new.width, new.height = w, int(h * w / new.width)
        new.left = left + (w - new.width) // 2 if plate else left
        new.name = 'Client logo'
    else:
        tb = sl.shapes.add_textbox(left, top, w, h)
        tb.name = 'Client wordmark'
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        if plate:
            lines_, size, spc = WMP, 12, '150'
        else:
            lines_, size, spc = WM, 14, '150'
            tb.width = Inches(4.9); tb.top = top - Inches(0.25)
        for i, txt in enumerate(lines_):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = PP_ALIGN.CENTER if plate else PP_ALIGN.LEFT
            r = para.add_run(); r.text = txt
            r.font.name = 'Cambria'; r.font.size = Pt(size); r.font.bold = True
            r.font.color.rgb = NAVY
            rPr = r._r.get_or_add_rPr(); rPr.set('spc', spc)
        new = tb
    parent.remove(new._element); parent.insert(pos, new._element)

# ---------- slide 1 cover
slot(1, 2, 'cover.date'); slot(1, 5, 'cover.client_name'); slot(1, 6, 'cover.client_addr')
replace_logo(1, 4)
# ---------- slide 2 intro
slot(2, 1, 'intro.title'); slot(2, 2, 'intro.addr'); slot(2, 3, 'intro.body')
S[1][2].width = Inches(3.7)
for _para in S[1][2].text_frame.paragraphs:
    for _r in _para.runs: _r.font.size = Pt(11)
# ---------- slide 3 contents
for n, (t, d) in enumerate([(4, 5), (8, 9), (12, 13), (16, 17), (20, 21), (24, 25)], 1):
    slot(3, t, f'contents.title{n}'); slot(3, d, f'contents.desc{n}')
# ---------- slide 4 who we are
slot(4, 2, 'who.intro')
for n, (h, b) in enumerate([(5, 6), (9, 10), (13, 14)], 1):
    slot(4, h, f'who.item{n}.head'); slot(4, b, f'who.item{n}.body')
for n, (h, b) in enumerate([(17, 18), (19, 20), (21, 22), (23, 24)], 1):
    slot(4, h, f'who.staff{n}.head'); slot(4, b, f'who.staff{n}.body')
# ---------- slide 5 team
slot(5, 2, 'team.intro')
for r in range(1, 5):
    for c in range(4): cell(5, 3, r, c, f'team.r{r}c{c}')
slot(5, 4, 'team.foot')
# ---------- slide 6 understanding
slot(6, 1, 'und.title')
slot(6, 4, 'und.company', spacer_idx=1, body_idx=2)
for n, i in enumerate([8, 11, 14, 17, 20], 1): slot(6, i, f'und.need{n}')
for n, (c, t) in enumerate([(23, 24), (26, 27), (29, 30), (32, 33), (35, 36)], 1):
    slot(6, c, f'und.date{n}.chip'); slot(6, t, f'und.date{n}.text')
slot(6, 37, 'und.foot')
# ---------- slide 7 services
slot(7, 1, 'svc.title')
for n, (h, b) in enumerate([(3, 4), (6, 7), (9, 10), (12, 13), (15, 16)], 1):
    slot(7, h, f'svc.card{n}.head'); slot(7, b, f'svc.card{n}.body')
slot(7, 17, 'svc.foot')
# ---------- slide 8 work together
for n, (h, b) in enumerate([(3, 4), (6, 7), (9, 10)], 1):
    slot(8, h, f'work.card{n}.head'); slot(8, b, f'work.card{n}.body')
slot(8, 12, 'work.cadence'); slot(8, 13, 'work.foot')
# ---------- slide 9 timeline
slot(9, 1, 'tl.title')
if 'tl.months' in C:
    for i, m in zip((3, 5, 7, 9, 11, 13), C['tl.months']): set_lines(S[8][i].text_frame, [m])
X0, MW = 4.50, 1.37
if 'tl.rows' in C:
    labels, bars = [15, 17, 20, 22, 25, 27, 30, 32, 35, 37], [16, 18, 21, 23, 26, 28, 31, 33, 36, 38]
    for li, bi, line in zip(labels, bars, C['tl.rows']):
        lab, st, en, col = [x.strip() for x in line.split('|')]
        st, en = float(st), float(en)
        set_lines(S[8][li].text_frame, [lab])
        bar = S[8][bi]
        bar.left, bar.width = Inches(X0 + st * MW), Inches(max(0.12, (en - st) * MW))
        bar.fill.solid(); bar.fill.fore_color.rgb = RED if col.lower().startswith('red') else NAVY
if 'tl.checkpoints' in C:
    for (di, ti), line in zip([(45, 46), (47, 48), (49, 50), (51, 52), (53, 54)], C['tl.checkpoints']):
        lab, pos = [x.strip() for x in line.split('|')]; pos = float(pos)
        set_lines(S[8][ti].text_frame, [lab])
        S[8][di].left = Inches(X0 + pos * MW - 0.11)
        S[8][ti].left = Inches(X0 + pos * MW - 0.70)
slot(9, 56, 'tl.foot')
# ---------- slide 10 90 days
for n, (r, nm, a, o) in enumerate([(4, 5, 7, 9), (12, 13, 15, 17), (20, 21, 23, 25)], 1):
    slot(10, r, f'p90.phase{n}.range'); slot(10, nm, f'p90.phase{n}.name')
    slot(10, a, f'p90.phase{n}.activities'); slot(10, o, f'p90.phase{n}.outcomes')
slot(10, 27, 'p90.day91')
# ---------- slide 11 deliverables
for r in range(1, 7):
    for c in range(4): cell(11, 2, r, c, f'del.r{r}c{c}')
slot(11, 3, 'del.foot')
# ---------- slide 12 fees
slot(12, 2, 'fees.intro')
for r in range(1, 4):
    for c in range(3): cell(12, 3, r, c, f'fees.r{r}c{c}')
slot(12, 4, 'fees.note'); slot(12, 7, 'fees.billing'); slot(12, 9, 'fees.budget'); slot(12, 11, 'fees.notincl')
# ---------- slide 13 why / next
slot(13, 2, 'why.sub')
for n, (h, b) in enumerate([(5, 6), (9, 10), (13, 14), (17, 18)], 1):
    slot(13, h, f'why.item{n}.head'); slot(13, b, f'why.item{n}.body')
for n, (h, b) in enumerate([(22, 23), (25, 26), (28, 29), (31, 32)], 1):
    slot(13, h, f'next.step{n}.head'); slot(13, b, f'next.step{n}.body')
set_lines(S[12][33].text_frame, ['Steven Nikolov, Principal  |  steven@onyxcfo.com  |  direct 480-999-5509'])
# ---------- slide 14 closing
slot(14, 1, 'close.thanks'); slot(14, 5, 'close.prepared')
replace_logo(14, 4, plate=True)

# ---------- layout footer, notes, properties
for layout in p.slide_layouts:
    for sh in layout.shapes:
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    if 'AJAC' in r.text:
                        r.text = f'Proposal for {CLIENT}  |  Confidential'
p.slides[0].notes_slide.notes_text_frame.text = 'Cover. Proposal accompanies the ACHS engagement letter.'
p.core_properties.title = f'ONYX Proposal: {CLIENT}'
p.core_properties.subject = 'Proposal for Accounting, Controller and CFO Services'
p.save(out)
print('saved', out)
